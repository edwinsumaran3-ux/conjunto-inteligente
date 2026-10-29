from __future__ import annotations
from io import BytesIO
from typing import Any
REQUIRED_SHEETS=("Component","System","Facility","Space","Type")
def parse_cobie_xlsx(content:bytes)->dict[str,Any]:
    if not content: raise ValueError("COBie XLSX is empty")
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(content),read_only=True,data_only=True); sheets=wb.sheetnames; missing=[s for s in REQUIRED_SHEETS if s not in sheets]; rows={}
    for s in sheets:
        it=wb[s].iter_rows(values_only=True)
        try: headers=next(it)
        except StopIteration: rows[s]=[]; continue
        headers=[str(x).strip() if x is not None else "" for x in headers]; out=[]
        for row in it:
            if any(x is not None and str(x).strip() for x in row): out.append({headers[i]:row[i] for i in range(min(len(headers),len(row))) if headers[i]})
        rows[s]=out
    comps=rows.get("Component",[]); return {"sheet_names":sheets,"missing_required_sheets":missing,"valid":not missing,"rows":rows,"component_count":len(comps),"components_missing_name":sum(1 for r in comps if not str(r.get("Name") or "").strip())}
